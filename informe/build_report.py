#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Genera Informe_Laboratorio5.docx (2 paginas A4) con python-docx.
Uso: python informe/build_report.py   (requiere python-docx; el PDF se obtiene con Word o LibreOffice)"""
import os
import re
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
EV = os.path.join(RAIZ, "evidencias")
SALIDA = os.path.join(RAIZ, "Informe_Laboratorio5.docx")

AZUL = RGBColor(0x1F, 0x4E, 0x9E)
GRIS = RGBColor(0x55, 0x5B, 0x66)
VERDE = RGBColor(0x1E, 0x7E, 0x4B)
ROJO = RGBColor(0xB0, 0x3A, 0x2E)
FUENTE = "Calibri"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.2)
sec.top_margin, sec.bottom_margin = Cm(1.0), Cm(0.9)
ANCHO = 21.0 - 2.4

st = doc.styles["Normal"]
st.font.name = FUENTE
st.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
st.font.size = Pt(8)
st.paragraph_format.space_after = Pt(1)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0


def runs(p, texto, negrita=False, tam=None, color=None, cursiva=False):
    """Formato en linea: **negrita**, `codigo` (Consolas), *cursiva*."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)", texto):
        if not tok:
            continue
        neg, cur, cod = negrita, cursiva, False
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, neg = tok[2:-2], True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            tok, cod = tok[1:-1], True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, cur = tok[1:-1], True
        r = p.add_run(tok)
        r.bold, r.italic = neg, cur
        if cod:
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if tam:
            r.font.size = Pt(tam - 0.7 if cod else tam)
        if color is not None:
            r.font.color.rgb = color
    return p


def parrafo(texto="", negrita=False, tam=None, color=None, alinear=None, antes=0, despues=1, contenedor=None, cursiva=False):
    p = (contenedor or doc).add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    if alinear:
        p.alignment = alinear
    if texto:
        runs(p, texto, negrita, tam, color, cursiva)
    return p


def titulo(texto):
    p = parrafo(texto, negrita=True, tam=9.5, color=AZUL, antes=3, despues=1)
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bt = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "1F4E9E")):
        bt.set(qn("w:" + k), v)
    b.append(bt)
    pPr.append(b)
    return p


def vinneta(texto, contenedor=None, tam=None):
    p = parrafo("", contenedor=contenedor, despues=0.5)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    runs(p, "• " + texto, tam=tam)
    return p


def sombra(celda, hexcolor):
    tcPr = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexcolor)
    tcPr.append(sh)


def margenes(celda, arriba=15, abajo=15, izq=40, der=40):
    tcPr = celda._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", arriba), ("left", izq), ("bottom", abajo), ("right", der)):
        e = OxmlElement("w:" + k)
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcPr.append(m)


def bordes(tabla, color="B8BEC9", sz="4"):
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + k)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), sz)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tabla._tbl.tblPr.append(b)


def sin_bordes(t):
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + k)
        e.set(qn("w:val"), "nil")
        b.append(e)
    t._tbl.tblPr.append(b)


def tabla(filas, anchos, tam=7, zebra=True):
    t = doc.add_table(rows=len(filas), cols=len(anchos))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bordes(t)
    for i, fila in enumerate(filas):
        for j, txt in enumerate(fila):
            c = t.cell(i, j)
            c.width = Cm(anchos[j])
            margenes(c)
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            es_cab = i == 0
            runs(p, txt, negrita=es_cab or j == 0, tam=tam, color=RGBColor(255, 255, 255) if es_cab else None)
            if es_cab:
                sombra(c, "1F4E9E")
            elif zebra and i % 2 == 0:
                sombra(c, "EEF3FB")
    return t


def imagen(celda, ruta, ancho_cm, pie=None, tam_pie=6.5):
    p = celda.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(ruta, width=Cm(ancho_cm))
    if pie:
        q = celda.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(0)
        runs(q, pie, tam=tam_pie, color=GRIS, cursiva=True)


def recorte(nombre, caja):
    """Recorta la zona relevante de una captura para que el texto sea legible en el informe.
    El archivo original de evidencias/ no se modifica; el recorte va a una carpeta temporal."""
    import tempfile
    from PIL import Image
    ruta = os.path.join(EV, nombre)
    if not caja:
        return ruta
    d = os.path.join(tempfile.gettempdir(), "lab5_recortes")
    os.makedirs(d, exist_ok=True)
    out = os.path.join(d, os.path.splitext(nombre)[0] + "_recorte.png")
    Image.open(ruta).crop(caja).save(out)
    return out


def cuadricula(items, cols, ancho_total, tam_pie=7):
    filas = (len(items) + cols - 1) // cols
    tt = doc.add_table(rows=filas, cols=cols)
    tt.autofit = False
    sin_bordes(tt)
    w = ancho_total / cols
    for k, it in enumerate(items):
        r, pie = it[0], it[1]
        caja = it[2] if len(it) > 2 else None
        c = tt.cell(k // cols, k % cols)
        c.width = Cm(w)
        margenes(c, 10, 10, 15, 15)
        imagen(c, recorte(r, caja), w - 0.3, pie, tam_pie=tam_pie)
    return tt


import json
from datetime import datetime, timedelta

with open(os.path.join(RAIZ, "mediciones", "20_resumen_toma.json"), encoding="utf-8") as fh:
    TOMA = json.load(fh)
COMPLETA = TOMA["duracion_dias"] >= 3.98
GRAFANA = "http://172.183.133.93:3000"
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


def bog(s):
    """Marca UTC de InfluxDB -> 'dd-mes hh:mm:ss' en hora de Bogota."""
    base = s.rstrip("Z").replace("+00:00", "").partition(".")[0]
    d = datetime.strptime(base, "%Y-%m-%dT%H:%M:%S") - timedelta(hours=5)
    return "%02d-%s %s" % (d.day, MESES[d.month - 1], d.strftime("%H:%M:%S"))


def miles(x):
    return "{:,}".format(x).replace(",", " ")


# ------------------------------------------------------------------ PAGINA 1
parrafo("Laboratorio 5 — Visualización + reglas/alertas: MQTT, AMQP y CoAP con Grafana en una VM", negrita=True, tam=12.5, color=AZUL, despues=0)
parrafo("IoT + Cloud + Sistemas Distribuidos · UNAB · Escenario UNAB-Ambiental · Plataforma propia (sin IoT Central) · "
        "Repositorio: https://github.com/Daniverd15/Laboratorio-5 · Dashboard en vivo (solo lectura): " + GRAFANA, tam=7.5, color=GRIS, despues=0)
parrafo("Integrantes: Daniel Villamizar · Tomás Urieles · David Guerrero", negrita=True, tam=8.5, despues=2)

titulo("1. Qué se montó: protocolo → TSDB → Grafana")
parrafo("En una VM de Azure (Ubuntu 22.04, B2ats_v2: 2 vCPU y **1 GB de RAM**) se levantó con un solo `docker-compose.yml` el stack de la plataforma propia: "
        "9 contenedores. Tres publicadores envían cada 15 s las 3 variables (`temperature`, `humidity`, `illuminance`) y la marca de tiempo real del nodo, "
        "**cada uno por un protocolo**; un ingestor las escribe en **InfluxDB 2** y **Grafana OSS**, que solo lee la TSDB, las muestra y evalúa las alertas (Fig. 1).", tam=8.5)
t = doc.add_table(rows=1, cols=1)
sin_bordes(t)
imagen(t.cell(0, 0), os.path.join(EV, "00-arquitectura.png"), 12.0,
       "**Fig. 1.** Diagrama del stack. El ingestor solo confirma al emisor (PUBACK, `basic.ack`, `2.04`) cuando InfluxDB ya aceptó el punto.")

titulo("2. Protocolo elegido y justificación")
parrafo("El enunciado deja elegir uno y permite usar más: **se implementaron los tres**, un nodo por protocolo, sobre la misma TSDB. "
        "**Protocolo principal declarado: MQTT**; AMQP es la versión **0-9-1** (RabbitMQ).", tam=8.5)
tabla([
    ["Nodo", "Protocolo y broker", "Cómo publica", "Garantía y qué pasa si el ingestor cae", "Por qué (y dónde no)"],
    ["nodo-mqtt", "MQTT 3.1.1 · Mosquitto 2.0 · TCP 1883", "tópico `unab/ambiental/<nodo>/telemetria`, usuario/clave, LWT",
     "QoS 1 (PUBACK). El broker guarda la sesión persistente del ingestor y reentrega.",
     "**Principal.** Cabecera de 2 B, una conexión saliente; lo hablan el ESP32 y la Raspberry Pi. No encola por sí solo sin sesión persistente."],
    ["nodo-amqp", "**AMQP 0-9-1** · RabbitMQ 3.13 · TCP 5672", "exchange topic `unab.ambiental`, rk `telemetria.<nodo>` → cola durable",
     "Mensaje persistente + *publisher confirm* + `mandatory`. La cola durable acumula.",
     "Para integrar servicios de la plataforma (colas, enrutamiento). Pesado para un microcontrolador."],
    ["nodo-coap", "CoAP RFC 7252 · aiocoap · UDP 5683", "`CON POST /telemetria` (JSON, Content-Format 50)",
     "ACK `2.04 Changed`. Sin broker: el nodo reintenta desde su cola local.",
     "El más liviano: **131 B de UDP por muestra, ACK de 6 B**. Sin buffer en el servidor y difícil tras NAT."],
], [1.5, 3.4, 4.0, 4.6, 5.1], tam=7)
parrafo("Medido en la VM: retardo dispositivo → TSDB **p50 = 1–2 ms** en los tres (todo en la red de Docker); la diferencia práctica no es la velocidad sino "
        "**quién guarda el mensaje cuando algo falla** (sección 4). Las claves van en `.env`, fuera del repositorio.", tam=8, antes=1)

titulo("3. Grafana aparte y alerta de umbral")
parrafo("Grafana 11.6 corre en su propio contenedor, con un **datasource de solo lectura** hacia InfluxDB (Flux) y todo aprovisionado desde archivos "
        "(datasource, dashboard y reglas). El dashboard tiene título, una línea de **qué mira el operador**, *gauges* con unidad y rango del escenario "
        "(temperatura 10–45 °C, humedad 0–100 %, iluminancia 0–100 000 lux), la serie de cada variable sobre **los últimos 4 días** y el umbral dibujado.", tam=8.5)
tabla([
    ["Regla", "Condición (umbral)", "Ventana", "Sostenida", "Acción"],
    ["**Temperatura alta** (la del laboratorio)", "promedio por nodo **> 32 °C**", "últimos 2 min, evaluada cada 30 s", "1 min",
     "**Webhook** → servicio `alertas`: línea JSON en `alertas.log` + log + republicación por MQTT en `unab/ambiental/alertas/#`. Avisa al disparar y al resolver."],
    ["Humedad baja", "promedio por nodo < 35 %HR", "últimos 2 min", "1 min", "igual"],
    ["Nodo sin datos", "último punto hace > 90 s", "mira 30 min", "30 s", "igual (vigila la toma de 4 días)"],
], [3.6, 3.4, 3.4, 1.4, 6.8], tam=7)
parrafo("**Disparo (13-oct-2026, hora de Bogotá):** `./scripts/forzar.sh nodo-mqtt temperature 41 300` hace que **el publicador** MQTT envíe 41 °C → "
        "forzado 00:43:12 → **Pending 00:44:46** → **Firing 00:45:46** → webhook 00:45:56 → fin del forzado 00:48:12 → **Normal 00:49:16** (Figs. 3–5). "
        "Los otros dos nodos siguieron en Normal: la regla es por nodo.", tam=8.5, antes=1)

titulo("4. Ventana de la toma de 4 días: inicio, fin y huecos")
n = TOMA["nodos"]
filas = [["Nodo", "Protocolo", "Primer punto (Bogotá)", "Puntos guardados", "Esperados (1 cada 15 s)", "Completitud", "Huecos > 30 s"]]
for k in sorted(n):
    d = n[k]
    filas.append([k, d["protocolo"].upper(), bog(d["inicio"]), miles(d["puntos"]), miles(d["esperados"]),
                  "%.2f %%" % d["completitud_pct"], "%d (%d s)" % (d["huecos"], d["segundos_sin_datos"])])
tabla(filas, [2.0, 1.8, 3.4, 2.8, 3.4, 2.4, 2.8], tam=7)
ini, fin = TOMA["inicio_bogota"], TOMA["fin_bogota"]
if COMPLETA:
    estado = ("**Toma completa:** inicio **%s**, fin **%s** (hora de Bogotá; ahí se detuvieron los publicadores) = **%.2f días** continuos. "
              "La ventana de 4 días que se entrega es **13-oct 00:35 → 17-oct 00:35** (Fig. 2); la toma siguió corriendo hasta el fin indicado." % (ini, fin, TOMA["duracion_dias"]))
else:
    estado = ("**Inicio: %s** (hora de Bogotá). **Fin previsto: 2026-10-17 00:35.** **Fecha de corte de este informe: %s** (%.2f h de toma); "
              "la toma sigue corriendo en la VM y la captura de la ventana completa se adjunta al entregar." % (ini, fin, TOMA["duracion_h"]))
parrafo(estado, tam=8.5, antes=1)
vinneta("**Hueco provocado:** `nodo-amqp`, 13-oct 00:50:34 → 00:53:49 (**195 s**): se detuvo el publicador a propósito; la regla *Nodo sin datos* pasó a Firing a los 129 s "
        "y el panel *Muestras por minuto* cayó a 0 (Fig. 8).", tam=8)
vinneta("**Sin hueco:** con el ingestor detenido 73 s no se perdió nada: Mosquitto retuvo la sesión QoS 1, RabbitMQ acumuló 4 mensajes en la cola durable y el nodo CoAP "
        "reintentó; al volver, todo se escribió con su marca de tiempo original (Fig. 8).", tam=8)
otros = [h for h in TOMA["huecos"] if not (h["nodo"] == "nodo-amqp" and h["desde"].startswith("2026-10-13T05:50"))]
if otros:
    vinneta("**Otros huecos (%d):** " % len(otros) + "; ".join("%s %s (%d s)" % (h["nodo"], bog(h["desde"]), h["segundos"]) for h in otros[:8]) + ".", tam=8)
else:
    vinneta("**No hay más huecos** a la fecha de corte. El pico de 41 °C de `nodo-mqtt` (13-oct 00:43–00:48) es el valor forzado, marcado `forzado=true` en la TSDB.", tam=8)

titulo("5. Qué falló al instalar en la VM")
for txt in [
    "**RAM:** la VM tiene 891 MB para 9 contenedores → 2 GB de swap, `mem_limit` por servicio y `vm_memory_high_watermark` de RabbitMQ en 256 MiB (uso real ≈ 430 MB + 330 MB de swap).",
    "**Grafana no cargaba el dashboard** (`Dashboard cannot be changed to a folder`): la carpeta aprovisionada tenía el mismo UID que el dashboard; se cambió `folderUid` y hubo que borrar el volumen `grafana-data`.",
    "**El publicador MQTT se reiniciaba** (`RuntimeError: Message publish failed`): paho falla si se publica antes del CONNACK → ahora publica solo si `is_connected()` y la muestra espera en cola.",
    "**Menores:** primer POST CoAP sin ACK porque el nodo arrancó antes que el servidor (lo resolvió la cola local); reloj sin sincronizar tras reactivar la VM (`timedatectl set-ntp true`); "
    "`join` de Flux sobre tablas vacías; `fromisoformat` de Python 3.10; CRLF en los `.sh` editados desde Windows (`.gitattributes`).",
]:
    vinneta(txt, tam=8)
# ------------------------------------------------------------------ PAGINA 2
doc.add_page_break()
titulo("6. Evidencia (capturas y logs reales de la VM; el resto está en /evidencias y /mediciones del repositorio)")
t = doc.add_table(rows=1, cols=1)
sin_bordes(t)
if COMPLETA:
    pie2 = "**Fig. 2.** Dashboard de Grafana con la **ventana de 4 días** (13-oct 00:35 → 17-oct 00:35): ciclo día/noche de las 3 variables en los 3 nodos, umbrales en rojo y el pico forzado de la alerta el primer día."
else:
    pie2 = ("**Fig. 2.** Dashboard de Grafana con la toma desde su inicio hasta la fecha de corte: *gauges* con unidad y rango, series de las 3 variables con el umbral, "
            "y el pico forzado de 41 °C con las marcas Pending (amarillo), Firing (rojo) y Normal (verde).")
imagen(t.cell(0, 0), recorte("20-dashboard-toma.png", (0, 440, 1700, 1440)), 15.0, pie2, tam_pie=7)
cuadricula([
    ("10-alerta-dashboard-firing.png", "**Fig. 3.** Alerta **disparada**: gauge de `nodo-mqtt` en 41 °C y serie sobre el umbral (corazón roto = Firing).", (0, 215, 1700, 690)),
    ("14-alerta-historial-estados.png", "**Fig. 4.** Historial de la regla en Grafana: Pending → Alerting → Normal, con el valor evaluado (B = promedio).", (320, 60, 1480, 610)),
], 2, ANCHO)
cuadricula([
    ("16-alerta-disparo-log.png", "**Fig. 5.** Forzado desde el publicador, cambios de estado y **acción**: el webhook recibe FIRING y RESOLVED y republica por MQTT."),
    ("02-publicadores-mqtt-amqp-coap.png", "**Fig. 6.** Las 3 variables por **MQTT, AMQP 0-9-1 y CoAP**, y el ingestor escribiendo en InfluxDB (204).", (0, 0, 1104, 560)),
], 2, ANCHO)
cuadricula([
    ("03-consulta-tsdb.png", "**Fig. 7.** **Consulta** Flux directa a la TSDB: puntos crudos con la marca de tiempo del dispositivo."),
    ("17-hueco-y-resiliencia.png", "**Fig. 8.** Hueco provocado (195 s) detectado y medido; ingestor caído 73 s **sin pérdida**.", (0, 0, 1248, 560)),
], 2, ANCHO)
parrafo("Las horas de Grafana están en hora de Bogotá (UTC−5); las de los logs y de InfluxDB, en UTC (05:45Z = 00:45 de Bogotá). "
        "Stack reproducible: `sudo ./scripts/instalar_vm.sh && cp .env.example .env && sudo docker compose up -d --build` (README).", tam=7.5, color=GRIS, antes=2)

parrafo("**Límites declarados:** sensores simulados (modelo día/noche); sin TLS/DTLS dentro de la red de Docker (solo Grafana sale a Internet); la acción es un webhook propio, no correo; una sola VM.",
        tam=7.5, color=GRIS, antes=1)

doc.save(SALIDA)
print("OK", SALIDA)
